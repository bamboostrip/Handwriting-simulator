"""docx 表格识别导入的渲染管线测试（对齐 Rust 0.5.0 的 doc_render.rs）。

覆盖：
- 换行签名合并 ``merge_wrap_line_boxes``（零重叠跨行合并 / 不同色 / 行尾非尾 /
  行首非首 / 垂直间隙过大 / 空字符流 / 1px 相触 / 容差 / 禁止链式吞并）；
- 字符级擦白 ``erase_char_boxes`` 保留表格线；
- ``process_tables_on_page`` 端到端（自造 ``ExtractedChar`` 网格，无需 Word/PDF）；
- 表内高亮框不整框擦白（保表格线）、表内标签不生成独立区域；
- ``TextRegion.table_index`` 序列化向后兼容与 ``label()`` 前缀；
- ``_word_com_script`` 的 Quit 容错与显式退出码；
- ``pdf_to_images_with_regions(body_stream=...)`` 全链路（手工构造最小文本 PDF）。

所有用例离线可跑：docx body 流用 dataclass 假流（不经文件系统），PDF 由
最小 PDF 字节手工拼装，不依赖 Microsoft Word / LibreOffice。
"""

from __future__ import annotations

import copy
from dataclasses import asdict
from pathlib import Path

import numpy as np
import pytest
from PIL import Image

from handwritesim.core.doc_render import (
    BoundingBox,
    ExtractedChar,
    _partition_highlight_boxes,
    _rect_contains_box,
    _word_com_script,
    erase_char_boxes,
    erase_highlight_boxes,
    extract_pdf_page_tags,
    merge_wrap_line_boxes,
    pdf_to_images_with_regions,
    process_tables_on_page,
)
from handwritesim.core.models import HandwritingParams, TextRegion


# ---------------------------------------------------------------------------
# 构造工具
# ---------------------------------------------------------------------------
def ch_at(ch: str, x: float, y: float, w: float, h: float) -> ExtractedChar:
    """构造一个字符对象：占据 [x, x+w) × [y, y+h)。"""
    return ExtractedChar(ch, x, y, x + w, y + h, 10.5, 10.5)


def grid_chars(text: str, x0: float, y0: float, size: float = 20.0) -> list[ExtractedChar]:
    """按等宽网格构造一行字符（表格单元格的印刷字）。"""
    return [
        ExtractedChar(
            ch, x0 + i * size, y0, x0 + (i + 1) * size, y0 + size, 10.5, 10.0,
        )
        for i, ch in enumerate(text)
    ]


def _cell_run(text: str, highlight: str | None = None):
    from handwritesim.core.docx_io import CellRun

    return CellRun(text=text, highlight=highlight)


def _parsed_cell(
    text: str | None = None,
    *,
    runs=None,
    col_start: int = 0,
    grid_span: int = 1,
    jc: str | None = None,
    sz_half_pt: int | None = None,
    v_merge=None,
):
    """构造 docx 解析层的单元格（延迟导入：表格结构模块由并行任务提供）。"""
    from handwritesim.core.docx_io import CellRun, CellVMerge, ParsedCell

    return ParsedCell(
        runs=list(runs) if runs is not None else [CellRun(text=text or "")],
        jc=jc,
        run_sz_half_pt=sz_half_pt,
        grid_span=grid_span,
        v_merge=CellVMerge.NONE if v_merge is None else v_merge,
        col_start=col_start,
    )


def _stream_of(rows: list[list]) -> tuple[list, list]:
    """rows -> (body 流, 表格视图)：走真实的 build_table_views，不造视图。"""
    from handwritesim.core.docx_io import ParsedRow, ParsedTable
    from handwritesim.core.docx_tables import build_table_views

    table = ParsedTable(
        grid_cols=max(len(r) for r in rows),
        rows=[ParsedRow(cells=list(r)) for r in rows],
    )
    stream = [table]
    return stream, build_table_views(stream)


def _blank_arr(w: int = 600, h: int = 400) -> np.ndarray:
    return np.full((h, w, 3), 255, dtype=np.uint8)


# ---------------------------------------------------------------------------
# 换行签名合并（merge_wrap_line_boxes）
# ---------------------------------------------------------------------------
# 复现几何（200DPI 真实 Word 链路实测）：句子从行中开始、下一行行首附近结束，
# 行尾条纹 (1315,1592)-(1403,1629) 与行首条纹 (250,1629)-(979,1666) 水平零重叠，
# 规则 3 判不合并 → 一句话被上下切成两个区域。
# 签名 = 上行条纹是「行尾」（同行右侧无字符）、下行条纹是「行首」（同行左侧无字符）。
def _wrap_scene_chars() -> list[ExtractedChar]:
    """己场景的字符流：行1 = 前缀(250..1310) + 句尾高亮(1315..1403)；
    行2 = 句首高亮(250..979) + 行2 后续正文(984..1390)。"""
    chars: list[ExtractedChar] = []
    for i in range(36):  # 行 1 前缀（未高亮）
        chars.append(ch_at("字", 250.0 + i * 29.4, 1592.0, 29.0, 37.0))
    for i in range(3):  # 行 1 句尾（高亮）
        chars.append(ch_at("字", 1315.0 + i * 29.4, 1592.0, 29.0, 37.0))
    for i in range(25):  # 行 2 句首（高亮）
        chars.append(ch_at("字", 250.0 + i * 29.4, 1629.0, 29.0, 37.0))
    for i in range(14):  # 行 2 后续未高亮正文
        chars.append(ch_at("字", 984.0 + i * 29.4, 1629.0, 29.0, 37.0))
    return chars


def _wrap_pair_boxes() -> list[BoundingBox]:
    return [
        BoundingBox(1315, 1592, 1403, 1629, "green"),
        BoundingBox(250, 1629, 979, 1666, "green"),
    ]


def test_wrap_tail_and_head_merge_into_one() -> None:
    merged = merge_wrap_line_boxes(_wrap_pair_boxes(), _wrap_scene_chars())
    assert len(merged) == 1, "行尾条纹+行首条纹应按换行签名合并为 1 个框"
    assert (merged[0].min_x, merged[0].max_x) == (250, 1403)
    assert (merged[0].min_y, merged[0].max_y) == (1592, 1666)
    assert merged[0].highlight == "green"


def test_wrap_merge_survives_reversed_input_order() -> None:
    boxes = [_wrap_pair_boxes()[1], _wrap_pair_boxes()[0]]
    merged = merge_wrap_line_boxes(boxes, _wrap_scene_chars())
    assert len(merged) == 1, "输入顺序不应影响换行签名合并"


def test_wrap_merge_respects_color() -> None:
    boxes = [
        BoundingBox(1315, 1592, 1403, 1629, "green"),
        BoundingBox(250, 1629, 979, 1666, "blue"),
    ]
    assert len(merge_wrap_line_boxes(boxes, _wrap_scene_chars())) == 2, \
        "不同颜色的行尾/行首条纹不得合并"


def test_wrap_merge_rejects_large_vertical_gap() -> None:
    # 行首条纹下移 200px（中间隔了未高亮内容），不构成同句换行
    boxes = [
        BoundingBox(1315, 1592, 1403, 1629, "green"),
        BoundingBox(250, 1830, 979, 1867, "green"),
    ]
    assert len(merge_wrap_line_boxes(boxes, _wrap_scene_chars())) == 2, \
        "垂直间隙过大不得按换行签名合并"


def test_wrap_merge_requires_tail_at_line_end() -> None:
    # 行 1 高亮条纹右侧还有正文（1410..1490）：句子没写到行尾，不是换行
    chars = _wrap_scene_chars()
    chars.append(ch_at("字", 1410.0, 1592.0, 29.0, 37.0))
    chars.append(ch_at("字", 1440.0, 1592.0, 29.0, 37.0))
    assert len(merge_wrap_line_boxes(_wrap_pair_boxes(), chars)) == 2, \
        "条纹不在行尾时不得合并"


def test_wrap_merge_requires_head_at_line_start() -> None:
    # 行 2 高亮条纹左侧还有正文（180..245）：句子不是从行首开始
    chars = _wrap_scene_chars()
    chars.append(ch_at("字", 180.0, 1629.0, 29.0, 37.0))
    chars.append(ch_at("字", 210.0, 1629.0, 29.0, 37.0))
    assert len(merge_wrap_line_boxes(_wrap_pair_boxes(), chars)) == 2, \
        "条纹不在行首时不得合并"


def test_wrap_merge_with_empty_chars_keeps_boxes() -> None:
    # 无字符流时行尾/行首判定空真，必须整体退化为不合并（防调用方漏守卫）
    assert len(merge_wrap_line_boxes(_wrap_pair_boxes(), [])) == 2, "空字符流不得合并"


def test_wrap_merge_single_box_and_empty_input() -> None:
    assert merge_wrap_line_boxes([], _wrap_scene_chars()) == []
    one = [BoundingBox(1315, 1592, 1403, 1629, "green")]
    assert merge_wrap_line_boxes(one, _wrap_scene_chars()) == one


def test_wrap_merge_allows_touching_stripes() -> None:
    # 相邻行条纹常在边界 1px 相触（gap_y 为负），应正常合并
    chars = _wrap_scene_chars()
    for c in chars:
        if c.min_y == 1629.0:
            c.min_y, c.max_y = 1628.0, 1665.0
    boxes = [
        BoundingBox(1315, 1592, 1403, 1629, "green"),
        BoundingBox(250, 1628, 979, 1665, "green"),
    ]
    assert len(merge_wrap_line_boxes(boxes, chars)) == 1, "1px 相触的换行对应合并"


def test_wrap_merge_tolerates_char_within_tol_right_of_tail() -> None:
    # 行尾紧邻的窄字符（如一个空格，落在半个框高容差内）不应阻断合并
    chars = _wrap_scene_chars()
    chars.append(ch_at("字", 1405.0, 1592.0, 14.0, 37.0))  # 距条纹右缘 2px 的窄字符
    assert len(merge_wrap_line_boxes(_wrap_pair_boxes(), chars)) == 1, \
        "容差内的紧邻字符不阻断换行合并"


def test_wrap_merge_does_not_chain_through_unrelated_blob_above() -> None:
    # 回归（集成实测踩坑）：上方无关多行大框与下方换行对间隙 108px，若允许
    # 联合框继续参与配对，其高度会让间隙阈值（1.5×较小框高）失真而整块吞并；
    # 单趟原始框配对必须保持「blob 独立 + 换行对合并」
    chars = _wrap_scene_chars()
    for i in range(38):
        chars.append(ch_at("字", 250.0 + i * 29.4, 1400.0, 29.0, 37.0))
    for i in range(38):
        chars.append(ch_at("字", 250.0 + i * 29.4, 1447.0, 29.0, 37.0))
    boxes = [
        BoundingBox(250, 1400, 1403, 1484, "green"),
        BoundingBox(1315, 1592, 1403, 1629, "green"),
        BoundingBox(250, 1629, 979, 1666, "green"),
    ]
    merged = merge_wrap_line_boxes(boxes, chars)
    assert len(merged) == 2, "blob 独立 + 换行对合并 = 2 框，不得链式吞并"
    wrap = next(b for b in merged if b.min_y == 1592)
    assert (wrap.min_y, wrap.max_y) == (1592, 1666)
    blob = next(b for b in merged if b.min_y == 1400)
    assert (blob.min_y, blob.max_y) == (1400, 1484)


# ---------------------------------------------------------------------------
# 字符级擦白（erase_char_boxes）
# ---------------------------------------------------------------------------
def test_erase_char_boxes_keeps_table_lines() -> None:
    # 300x200 白图；x=150 处竖直黑色表格线；(60..100)×(60..80) 有文字字形
    arr = _blank_arr(300, 200)
    arr[20:180, 150] = 0
    arr[60:80, 60:100] = 0
    chars = [ExtractedChar("字", 60.0, 60.0, 100.0, 80.0, 12.0, 0.0)]

    erase_char_boxes(arr, chars, [0], 2)

    assert (arr[70, 80] == 255).all(), "字形区域应被擦白"
    assert (arr[70, 150] == 0).all(), "表格线必须保留（整框擦白会把它抹掉）"


def test_erase_char_boxes_ignores_out_of_range_index() -> None:
    arr = _blank_arr(50, 50)
    arr[10:20, 10:20] = 0
    chars = [ExtractedChar("字", 10.0, 10.0, 20.0, 20.0, 12.0, 0.0)]
    # 越界下标只跳过该字符，不抛异常、不影响其余像素
    erase_char_boxes(arr, chars, [5], 2)
    assert (arr[15, 15] == 0).all()
    erase_char_boxes(arr, [], [0], 2)  # 空字符流不应崩
    assert (arr[15, 15] == 0).all()


def test_erase_char_boxes_clamps_to_image_bounds() -> None:
    # 包围盒越出图像右/下边界时钳制，不越界写
    arr = _blank_arr(40, 40)
    arr[0:40, 0:40] = 0
    chars = [ExtractedChar("字", 30.0, 30.0, 60.0, 60.0, 12.0, 0.0)]
    erase_char_boxes(arr, chars, [0], 2)
    assert (arr[38, 38] == 255).all()
    assert (arr[0, 0] == 0).all(), "只擦包围盒内，左上角保持不变"


def test_erase_char_boxes_uses_tight_span_to_spare_table_lines() -> None:
    """loose 框是「行高」，Word 表格行高紧贴边框时会盖住表格线。

    真实文档实测：``Table Grid`` 样式的行高≈行距，loose 框上下各压住一条
    边框，按 loose 擦会把整行边框擦成缺口。擦除必须用紧包围盒（≈字形高度）。
    """
    arr = _blank_arr(300, 200)
    arr[40, 20:280] = 0   # 上方表格线（y=40）
    arr[80, 20:280] = 0   # 下方表格线（y=80）
    # 文字 loose 框 y=40..80（正好压住两条线），紧框 y=52..68（字形本体）
    chars = [ExtractedChar("字", 60.0, 40.0, 100.0, 80.0, 12.0, 12.0, 52.0, 68.0)]

    erase_char_boxes(arr, chars, [0], 2)

    assert (arr[60, 80] == 255).all(), "字形区域应被擦白"
    assert (arr[40, 80] == 0).all(), "loose 框盖住的上表格线必须保留"
    assert (arr[80, 80] == 0).all(), "loose 框盖住的下表格线必须保留"


def test_erase_char_boxes_falls_back_to_loose_when_tight_missing() -> None:
    """未提供紧包围盒（旧数据/测试构造）时回退 loose，行为不变。"""
    arr = _blank_arr(120, 120)
    arr[40:60, 40:60] = 0
    chars = [ExtractedChar("字", 40.0, 40.0, 60.0, 60.0, 12.0, 10.0)]
    erase_char_boxes(arr, chars, [0], 2)
    assert (arr[50, 50] == 255).all()


def test_tag_erase_box_ignores_outlier_loose_boxes() -> None:
    """标签擦除框不能被标点的异常 loose 框撑大。

    真实文档：``{{打印:张三}}`` 里的 ``{ } :`` 在 Word/LibreOffice 导出的 PDF
    中字体矩阵异常（font_size 221pt），loose 框纵向撑到 277..448（跨整张表格），
    按 loose 擦标签会把该列上下四条表格线全部抹掉。修复：纵向用紧包围盒 +
    剔除高度离群字。
    """
    arr = _blank_arr(300, 200)
    arr[60, 0:300] = 0    # 上表格线
    arr[150, 0:300] = 0   # 下表格线
    # 标签 {{手写:答案}}：花括号 loose=60..150（跨两条线），tight=95..125
    chars: list[ExtractedChar] = [
        ExtractedChar("{", 10.0, 60.0, 22.0, 150.0, 221.0, 30.0, 95.0, 125.0),
        ExtractedChar("{", 22.0, 60.0, 34.0, 150.0, 221.0, 30.0, 95.0, 125.0),
        ExtractedChar("手", 34.0, 90.0, 64.0, 130.0, 20.0, 30.0, 95.0, 125.0),
        ExtractedChar("写", 64.0, 90.0, 94.0, 130.0, 20.0, 30.0, 95.0, 125.0),
        ExtractedChar(":", 94.0, 60.0, 102.0, 150.0, 221.0, 30.0, 100.0, 120.0),
        ExtractedChar("答", 102.0, 90.0, 132.0, 130.0, 20.0, 30.0, 95.0, 125.0),
        ExtractedChar("案", 132.0, 90.0, 162.0, 130.0, 20.0, 30.0, 95.0, 125.0),
        ExtractedChar("}", 162.0, 60.0, 174.0, 150.0, 221.0, 30.0, 95.0, 125.0),
        ExtractedChar("}", 174.0, 60.0, 186.0, 150.0, 221.0, 30.0, 95.0, 125.0),
    ]

    regions = extract_pdf_page_tags(chars, 0, 200, 300, 200, arr)

    assert len(regions) == 1
    region = regions[0]
    assert region.y >= 90 and region.y + region.h <= 132, (
        f"标签区域纵向应贴合字形（~95..125），实际 y={region.y} h={region.h}"
    )
    assert (arr[60, 100] == 0).all(), "上表格线不得被标签擦除框抹掉"
    assert (arr[150, 100] == 0).all(), "下表格线不得被标签擦除框抹掉"
    assert (arr[110, 40] == 255).all(), "标签文字本身要被擦白"


# ---------------------------------------------------------------------------
# 表内高亮框分流 / 表内标签排重
# ---------------------------------------------------------------------------
def test_rect_contains_box_checks_containment() -> None:
    from handwritesim.core.docx_tables import RectPx

    r = RectPx(10, 10, 100, 100)
    assert _rect_contains_box(r, BoundingBox(20, 20, 60, 60))
    assert not _rect_contains_box(r, BoundingBox(5, 20, 60, 60)), "越出左界"
    assert not _rect_contains_box(r, BoundingBox(20, 20, 120, 60)), "越出右界"


def test_table_highlight_box_is_not_erased_as_whole() -> None:
    """表内高亮框不整框擦白（保表格线），表外框仍整框擦白。"""
    from handwritesim.core.docx_tables import RectPx

    arr = _blank_arr(400, 240)
    # 表内：黄色高亮块 + 块内一条表格线片段
    arr[100:141, 40:141] = (255, 255, 0)
    arr[120, 90:101] = 0
    # 表外：黄色高亮块 + 块内一条黑色横线（模拟紧贴高亮的表格线/下划线）
    arr[100:141, 200:281] = (255, 255, 0)
    arr[120, 250:261] = 0

    table_outer = [RectPx(30, 60, 150, 120)]  # 覆盖表内块（x 30..180 / y 60..180）
    inside = BoundingBox(40, 100, 140, 140, "yellow")
    outside = BoundingBox(200, 100, 280, 140, "yellow")

    in_table, normal = _partition_highlight_boxes([inside, outside], table_outer)
    assert in_table == [inside], "完整落在表格外包围内的框应分流为表内"
    assert normal == [outside], "越出表格的框应留在表外"

    erase_highlight_boxes(arr, normal)

    assert (arr[120, 95] == 0).all(), "表内表格线必须保留（整框擦白会抹掉它）"
    assert (arr[110, 45] == 255).all(), "表内高亮像素仍被整页残留过滤清除（与 Rust 一致）"
    assert (arr[120, 255] == 255).all(), "表外框整框擦白：框内黑线被抹掉"


def _tag_chars() -> list[ExtractedChar]:
    """字符流中一个 ``{{姓名}}`` 标签（x 100..144 / y 50..70）。"""
    spans = [
        ("{", 100.0, 106.0), ("{", 106.0, 112.0),
        ("姓", 112.0, 122.0), ("名", 122.0, 132.0),
        ("}", 132.0, 138.0), ("}", 138.0, 144.0),
    ]
    return [
        ExtractedChar(ch, x0, 50.0, x1, 70.0, 10.5, 0.0)
        for ch, x0, x1 in spans
    ]


def test_extract_pdf_page_tags_excludes_table_cells() -> None:
    """与单元格矩形相交的标签：仍擦除标签框，但不生成独立区域。"""
    from handwritesim.core.docx_tables import RectPx

    chars = _tag_chars()
    arr = _blank_arr(400, 300)
    # 预涂黑标签区域，验证被单元格矩形覆盖时仍会擦除（模板字符不能残留底图）
    arr[48:72, 98:146] = 0
    exclude = [RectPx(90, 40, 70, 40)]
    regions = extract_pdf_page_tags(chars, 0, 200, 400, 300, arr, exclude_rects=exclude)
    assert regions == [], "表格单元格内的标签不应生成独立区域（避免与单元格区域重复）"
    assert (arr[60, 120] == 255).all(), "标签框仍必须被擦除"

    # 重叠但中心不在矩形内（标签框比单元格正文包围盒宽，如 {{打印:xxx}} 格）：
    # 同样接管——擦除且不生成区域
    overlap_only = [RectPx(135, 40, 70, 40)]
    arr2 = _blank_arr(400, 300)
    assert extract_pdf_page_tags(
        chars, 0, 200, 400, 300, arr2, exclude_rects=overlap_only
    ) == [], "重叠即接管，不应生成重复区域"

    # 不排除时能识别到（对照组）
    arr3 = _blank_arr(400, 300)
    regions3 = extract_pdf_page_tags(chars, 0, 200, 400, 300, arr3)
    assert len(regions3) == 1
    assert regions3[0].text == "姓名"


# ---------------------------------------------------------------------------
# process_tables_on_page 端到端（自造 ExtractedChar 网格）
# ---------------------------------------------------------------------------
def test_process_tables_builds_regions_and_erases() -> None:
    # 一张 1 行 2 列表：姓名 / 张三（黄色高亮 + 居中）
    stream, views = _stream_of([[
        _parsed_cell("姓名", col_start=0),
        _parsed_cell("张三", col_start=1, jc="center",
                     runs=[_cell_run("张三", highlight="yellow")]),
    ]])
    chars = grid_chars("姓名", 100.0, 100.0) + grid_chars("张三", 300.0, 100.0)
    arr = _blank_arr()
    # 印上「张三」字形（模拟底图上的印刷字）
    arr[100:120, 300:340] = 0

    color_map: dict[str, int] = {}
    next_role_id = [2]
    result = process_tables_on_page(
        stream, views, chars, 0, 1, 200, arr, color_map, next_role_id
    )

    assert result.next_start_element == 1, "2/2 定位完成 → 表格消费，流起点越过表格"
    assert len(result.regions) == 2
    r0, r1 = result.regions
    assert (r0.text, r0.table_index, r0.page) == ("姓名", 1, 1)
    assert r0.align == "left"
    assert r0.font_size > 0
    assert (r0.margin_left, r0.margin_right, r0.margin_top, r0.margin_bottom) == (4, 4, 2, 2)
    # 命中 bbox（pad=2）外扩 水平 6 / 竖直 4：(98,98,44,24) -> (92,94,56,32)
    assert (r0.x, r0.y, r0.w, r0.h) == (92, 94, 56, 32)
    # 字号经 _resolve_font_size_px 校准：全角紧包围盒 10pt × 200/72 → 28px
    assert r0.font_size == 28

    assert (r1.text, r1.highlight, r1.role_id) == ("张三", "yellow", 2)
    assert r1.align == "center", "w:jc=center → TextRegion.align='center'"
    assert (r1.x, r1.y, r1.w, r1.h) == (292, 94, 56, 32)
    assert color_map == {"yellow": 2} and next_role_id[0] == 3

    # 擦除后底图上的「张三」字形消失
    assert (arr[110, 310] == 255).all()
    # pad=2 之外的「姓名」位置本来没墨迹，仍为白
    assert (arr[110, 105] == 255).all()

    # 排除矩形与外包围
    assert len(result.cell_rects) == 2
    assert len(result.outer_rects) == 1
    outer = result.outer_rects[0]
    # 并集 x 98..342 / y 98..122 → 外扩 4
    assert (outer.x, outer.y, outer.w, outer.h) == (94, 94, 252, 32)


def _cell_run(text: str, highlight: str | None = None):
    from handwritesim.core.docx_io import CellRun

    return CellRun(text=text, highlight=highlight)


def test_process_tables_inferred_rect_not_expanded_and_font_fallback() -> None:
    """未命中格：矩形取行带×列带交集（已内缩），不再外扩；字号回退 w:sz。"""
    stream, views = _stream_of([
        [
            _parsed_cell("甲甲", col_start=0),
            _parsed_cell("乙乙", col_start=1),
        ],
        [
            _parsed_cell("丙丙", col_start=0),
            _parsed_cell("丁丁", col_start=1, sz_half_pt=21),
        ],
    ])
    chars = (
        grid_chars("甲甲", 100.0, 100.0)
        + grid_chars("乙乙", 300.0, 100.0)
        + grid_chars("丙丙", 100.0, 200.0)
        # 「丁丁」不在 PDF 字符流中（未命中格）
    )
    arr = _blank_arr()
    result = process_tables_on_page(
        stream, views, chars, 0, 1, 200, arr, {}, [2]
    )

    assert len(result.regions) == 4, "4 个非续行非空格都应生成区域"
    ding = next(r for r in result.regions if r.text == "丁丁")
    # 推断矩形 = (col 1 列带 x 298..342, row 1 行带 y 198..222) 内缩 3：
    # (301, 201, 38, 18)；命中格才会外扩，推断矩形保持原样
    assert (ding.x, ding.y, ding.w, ding.h) == (301, 201, 38, 18)
    # 无命中字符 → 字号回退 w:sz（21 半磅 = 10.5pt）→ 10.5×200/72 → 29
    assert ding.font_size == 29
    assert ding.role_id == 0 and ding.highlight is None


def test_process_tables_skips_gray_and_continue_cells() -> None:
    """续行格（vMerge=continue）与灰高亮格（skip_printed）不生成区域/不擦除。"""
    from handwritesim.core.docx_io import CellVMerge

    stream, views = _stream_of([
        [_parsed_cell("合并", col_start=0, v_merge=CellVMerge.RESTART)],
        [_parsed_cell("", col_start=0, v_merge=CellVMerge.CONTINUE)],
        [_parsed_cell(runs=[_cell_run("灰格", highlight="lightGray")], col_start=0)],
    ])
    chars = grid_chars("合并", 100.0, 100.0) + grid_chars("灰格", 100.0, 300.0)
    arr = _blank_arr()
    arr[300:320, 100:140] = 0  # 灰格的印刷底图

    result = process_tables_on_page(
        stream, views, chars, 0, 1, 200, arr, {}, [2]
    )

    assert [r.text for r in result.regions] == ["合并"], "续行格与灰格都不生成区域"
    assert (arr[310, 110] == 0).all(), "灰格保留印刷底图（不擦除）"
    assert (arr[110, 105] == 255).all(), "「合并」字形被擦除"


def test_process_tables_unlocated_table_yields_no_regions() -> None:
    """未定位的表整表跳过，但流起点仍指向表格元素（下一页重走）。"""
    from handwritesim.core.docx_io import BodyParagraph, ParsedRow, ParsedTable
    from handwritesim.core.docx_tables import build_table_views

    table = ParsedTable(grid_cols=2, rows=[
        ParsedRow(cells=[_parsed_cell("甲格", col_start=0), _parsed_cell("乙格", col_start=1)])
    ])
    stream = [BodyParagraph(text="第一页正文"), table]
    views = build_table_views(stream)
    chars = grid_chars("第一页正文", 10.0, 100.0, size=12.0)
    arr = _blank_arr()

    result = process_tables_on_page(
        stream, views, chars, 0, 1, 200, arr, {}, [2]
    )

    assert result.regions == [], "未定位表不得生成区域"
    assert result.cell_rects == []
    assert result.outer_rects == []
    assert result.next_start_element == 1, "流起点指向表格元素（下标 1），下一页重走"


def test_process_tables_shares_role_map_across_pages() -> None:
    """高亮色走共享 color_map：同色跨页复用角色 ID，新色 next_role_id 递增。"""
    stream1, views1 = _stream_of([[
        _parsed_cell(runs=[_cell_run("甲甲", highlight="yellow")], col_start=0),
    ]])
    stream2, views2 = _stream_of([[
        _parsed_cell(runs=[_cell_run("乙乙", highlight="blue")], col_start=0),
    ]])
    chars1 = grid_chars("甲甲", 100.0, 100.0)
    chars2 = grid_chars("乙乙", 100.0, 100.0)
    color_map: dict[str, int] = {}
    next_role_id = [2]

    r1 = process_tables_on_page(
        stream1, views1, chars1, 0, 1, 200, _blank_arr(), color_map, next_role_id
    )
    r2 = process_tables_on_page(
        stream2, views2, chars2, 0, 2, 200, _blank_arr(), color_map, next_role_id
    )
    # 回第 1 页重复的黄色复用同一角色
    r3 = process_tables_on_page(
        stream1, views1, chars1, 0, 3, 200, _blank_arr(), color_map, next_role_id
    )

    assert r1.regions[0].role_id == 2 and r1.regions[0].highlight == "yellow"
    assert r2.regions[0].role_id == 3 and r2.regions[0].highlight == "blue"
    assert r3.regions[0].role_id == 2, "同一高亮色跨页复用角色 ID"
    assert color_map == {"yellow": 2, "blue": 3}
    assert next_role_id[0] == 4
    assert r2.regions[0].page == 2


def test_process_tables_region_clamped_to_image() -> None:
    """单元格矩形越出图像边界时区域钳制到页内且不小于 1px。"""
    stream, views = _stream_of([[_parsed_cell("甲乙", col_start=0)]])
    # 字符贴右下角：命中 bbox 外扩后必然越界
    chars = grid_chars("甲乙", 570.0, 370.0)
    arr = _blank_arr(600, 400)

    result = process_tables_on_page(
        stream, views, chars, 0, 1, 200, arr, {}, [2]
    )

    assert len(result.regions) == 1
    r = result.regions[0]
    assert r.x + r.w <= 600 and r.y + r.h <= 400
    assert r.w >= 1 and r.h >= 1


# ---------------------------------------------------------------------------
# TextRegion.table_index：模型兼容性
# ---------------------------------------------------------------------------
def test_text_region_table_index_defaults_none_for_legacy_json() -> None:
    """旧预设/旧项目 JSON 无 table_index 字段 → 反序列化为 None。"""
    legacy = {"x": 1, "y": 2, "w": 10, "h": 5, "text": "答案", "page": 1}
    params = HandwritingParams.from_dict({"regions": [legacy]})
    assert params.regions is not None
    assert params.regions[0].table_index is None
    assert TextRegion().table_index is None

    # 新字段写入后往返保持（asdict → TextRegion(**d) 与项目反序列化同路径）
    r = TextRegion(x=1, y=2, w=10, h=5, text="答案", table_index=3)
    dumped = asdict(r)
    assert dumped["table_index"] == 3
    assert TextRegion(**dumped) == r
    # 新 dict 经 HandwritingParams.from_dict 保留 table_index
    params2 = HandwritingParams.from_dict({"regions": [dumped]})
    assert params2.regions is not None
    assert params2.regions[0].table_index == 3
    # 深拷贝（GUI 预览/整表批量应用路径）不丢字段
    assert copy.deepcopy(r).table_index == 3


def test_text_region_label_table_prefix() -> None:
    """表格单元格区域带 ``表N·`` 前缀，普通区域无前缀。"""
    plain = TextRegion(text="你好", page=2)
    assert plain.label(1) == "1. 手写 第2页 2字 (0,0 0×0)"
    cell = TextRegion(text="答案", table_index=2)
    assert cell.label(3) == "表2·3. 手写 2字 (0,0 0×0)"
    # 打印体 + 页码 + 表格前缀同时存在
    printed = TextRegion(text="张三", table_index=1, printed=True, page=5)
    assert printed.label(2) == "表1·2. 打印 第5页 2字 (0,0 0×0)"


# ---------------------------------------------------------------------------
# Word COM 脚本
# ---------------------------------------------------------------------------
def test_word_com_script_quit_is_tolerant_and_exit_is_explicit() -> None:
    """Quit 必须容错且退出码显式化。

    COM 拉起的 Word 在最后一个文档关闭后可能自行退出，Quit() 抛 RPC
    0x800706BE；而 powershell -Command 下即使异常被 try/catch 捕获，退出码
    仍会变 1 覆盖成功的转换，故脚本须以 PDF 存在性显式决定退出码。
    """
    script = _word_com_script(Path("a'b.docx"), Path("c'd.pdf"))
    assert "try { $word.Quit() } catch { }" in script, script
    assert "if (Test-Path 'c''d.pdf') { exit 0 } else { exit 1 }" in script, script
    assert "a''b.docx" in script and "c''d.pdf" in script, "单引号必须转义"


# ---------------------------------------------------------------------------
# 手工构造最小 PDF：全链路 pdf_to_images_with_regions(body_stream=...)
# ---------------------------------------------------------------------------
def _make_table_text_pdf(path) -> str:
    """最小 PDF：一行两块 Helvetica 文本 + 一个黄色高亮色块。

    - ``AB``  → 普通单元格（应生成单元格区域并擦除印刷字）；
    - ``{{hand:CD}}`` → 单元格内的标签（应擦除但不生成独立区域）；
    - 黄色矩形 → 表外高亮框（应生成普通区域，且排在表格区域之后）。
    """
    content = (
        b"BT /F1 12 Tf 72 700 Td (AB) Tj ET\n"
        b"BT /F1 12 Tf 200 700 Td ({{hand:CD}}) Tj ET\n"
        b"q 1 1 0 rg 300 600 60 20 re f Q"
    )
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()

    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(bytes(out))
    doc.save(str(path))
    doc.close()
    return str(path)


def test_pdf_pipeline_with_table_body_stream(tmp_path) -> None:
    """docx body 流驱动的全链路：表格区域 + 标签排重 + 表外高亮区域排序。"""
    pdf = _make_table_text_pdf(tmp_path / "table.pdf")
    # 单元格 1 = AB；单元格 2 原文含 {{hand:CD}} 标签（可见文本 CD）
    stream, views = _stream_of([[
        _parsed_cell("AB", col_start=0),
        _parsed_cell("{{hand:CD}}", col_start=1),
    ]])
    assert [c.text for c in views[0].rows[0].cells] == ["AB", "CD"], "标签应被剥离"

    paths, regions = pdf_to_images_with_regions(
        pdf, tmp_path / "out", dpi=72, body_stream=stream
    )

    assert len(paths) == 1
    table_regions = [r for r in regions if r.table_index == 1]
    assert [r.text for r in table_regions] == ["AB", "CD"], \
        "两个单元格各生成一个区域（标签不重复生成独立区域）"
    assert all(r.page == 1 and r.font_size > 0 for r in table_regions)
    # 表格区域先 extend，表外区域（黄色高亮框）在后
    assert [r.table_index for r in regions] == [1, 1, None], "表格区域必须排在表外区域之前"

    # 印刷字与标签字符都被擦白（单元格路径擦正文，标签路径擦模板字符）
    arr = np.array(Image.open(paths[0]).convert("RGB"))
    assert int((arr.min(axis=-1) < 128).sum()) == 0, "底图上不应残留黑色印刷字/标签"

    # 未传 body_stream 时行为与旧版一致：不产生表格区域
    paths2, regions2 = pdf_to_images_with_regions(pdf, tmp_path / "out2", dpi=72)
    assert all(r.table_index is None for r in regions2)
    assert len(paths2) == 1


def test_pdf_pipeline_table_regions_survive_body_stream_as_list(tmp_path) -> None:
    """body_stream 传元组/列表均可（document_to_page_images_with_regions 传列表）。"""
    pdf = _make_table_text_pdf(tmp_path / "table2.pdf")
    stream, _views = _stream_of([[_parsed_cell("AB", col_start=0)]])
    _paths, regions = pdf_to_images_with_regions(
        pdf, tmp_path / "out", dpi=72, body_stream=tuple(stream)
    )
    assert [r.table_index for r in regions if r.table_index is not None] == [1]


# ---------------------------------------------------------------------------
# 真实 docx 结构 → 视图 → 区域（不需要 Word/PDF）
# ---------------------------------------------------------------------------
def test_process_tables_from_real_docx(tmp_path) -> None:
    """python-docx 生成的真实 docx：load_body_stream → build_table_views →
    process_tables_on_page 全链路（表格结构解析走真实 OOXML 解析器）。"""
    from docx import Document

    from handwritesim.core.docx_io import load_body_stream
    from handwritesim.core.docx_tables import build_table_views

    doc = Document()
    doc.add_paragraph("前言")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "张三"
    path = tmp_path / "t.docx"
    doc.save(str(path))

    stream = load_body_stream(path)
    views = build_table_views(stream)
    assert len(views) == 1
    assert [c.text for c in views[0].rows[0].cells] == ["姓名", "张三"]

    chars = grid_chars("姓名", 100.0, 100.0) + grid_chars("张三", 300.0, 100.0)
    arr = _blank_arr()
    result = process_tables_on_page(stream, views, chars, 0, 1, 200, arr, {}, [2])

    assert [r.text for r in result.regions] == ["姓名", "张三"]
    assert all(r.table_index == 1 and r.page == 1 for r in result.regions)
    assert len(result.cell_rects) == 2 and len(result.outer_rects) == 1
    # 表格在流中位于「前言」之后（下标 1）→ 消费后起点越过表格
    assert result.next_start_element >= 2


def test_docx_entry_degrades_when_body_stream_parse_fails(tmp_path, monkeypatch) -> None:
    """docx 表格结构解析失败不得阻塞导入：退化为空流（纯底图，表格保持印刷）。"""
    import handwritesim.core.doc_render as doc_render
    import handwritesim.core.docx_io as docx_io

    def boom(_path):
        raise ValueError("解析失败")

    monkeypatch.setattr(docx_io, "load_body_stream", boom)
    monkeypatch.setattr(doc_render, "docx_to_pdf", lambda p, o: Path(p).with_suffix(".pdf"))
    captured: dict[str, object] = {}

    def fake_pdf(pdf_path, out_dir, dpi=200, prefix=None, body_stream=None):
        captured["body_stream"] = body_stream
        return [Path("page.png")], []

    monkeypatch.setattr(doc_render, "pdf_to_images_with_regions", fake_pdf)
    docx = tmp_path / "a.docx"
    docx.write_bytes(b"")
    paths, regions = doc_render.document_to_page_images_with_regions(
        docx, tmp_path / "out"
    )

    assert captured["body_stream"] == [], "解析失败应退化为空流"
    assert paths == [Path("page.png")] and regions == []

